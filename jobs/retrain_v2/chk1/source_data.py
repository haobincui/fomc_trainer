"""Fail-closed source-data primitives for the chk1 dataset rebuild.

This module intentionally stops before teacher generation or dataset release.
It provides four small, independently testable boundaries:

* immutable inventories for legacy inputs;
* a safe projection of the existing reference-free student-prompt schema onto
  the canonical ``meeting_date + atomic_topic`` sample key;
* point-in-time evidence and fact-card validation, including an adapter for the
  repository's canonical ALFRED/LOO ledger; and
* a tokenizer-independent 4,096-token prompt gate that never truncates.

Legacy prompts, responses, Minutes, rate decisions, and post-decision target
rates are audit inputs only.  They are deliberately absent from every
canonical sample and fact-card output produced here.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import unicodedata
from collections import defaultdict
from collections.abc import Callable, Iterable, Mapping, Sequence
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any


INVENTORY_SCHEMA_VERSION = "chk1-source-inventory-v1"
CANONICAL_SAMPLE_SCHEMA_VERSION = "chk1-canonical-analysis-sample-v1"
CANONICAL_BATCH_SCHEMA_VERSION = "chk1-canonical-analysis-batch-v1"
FACT_CARD_SCHEMA_VERSION = "chk1-point-in-time-fact-card-v1"
FACT_CARD_BUILD_SCHEMA_VERSION = "chk1-fact-card-build-v1"
PROMPT_BUDGET_SCHEMA_VERSION = "chk1-prompt-budget-v1"

PROMPT_TOKEN_LIMIT = 4096
ALFRED_SOURCE_INTERFACE = "alfred-graph-csv-v1"
ALFRED_AVAILABILITY_EVIDENCE = "alfred_vintage_snapshot"

# This is the schema emitted by build_analysis_student_prompts.py in the
# current repository.  Keeping the complete set here makes schema drift an
# explicit error instead of silently changing the canonical population.
STUDENT_PROMPT_REQUIRED_FIELDS = frozenset(
    {
        "sample_id",
        "split",
        "meeting_date",
        "section_name",
        "topic",
        "source_row_index",
        "rate_change",
        "current_rate",
        "prompt",
        "prompt_hash",
        "provided_data",
        "data_label",
        "reference_excerpt",
        "reference_source_status",
        "has_nonempty_table",
        "table_count",
        "data_table_count",
        "header_only_table_count",
        "prompt_length_chars",
        "prompt_length_words",
        "quality_flags",
        "missing_indicators",
        "source_files",
        "archived_response",
        "response_origin",
    }
)

ALLOWED_SPLITS = frozenset({"train", "eval", "test"})
FORBIDDEN_ATOMIC_TOPICS = frozenset(
    {"", "delete", "non-core", "non core", "other"}
)
ALLOWED_FACT_KINDS = frozenset(
    {
        "recent_observation",
        "latest",
        "change_3m",
        "change_6m",
        "change_12m",
        "change_1q",
        "change_2q",
        "change_4q",
        "mom",
        "qoq",
        "yoy",
        "trend",
        "turning_point",
        "extreme",
        "prior_target_range",
    }
)
DERIVED_FACT_KINDS = ALLOWED_FACT_KINDS - {
    "recent_observation",
    "latest",
    "prior_target_range",
}

_SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
_CLASSIFICATION_RE = re.compile(r"^[a-z][a-z0-9_]{0,63}$")
_MOJIBAKE_MARKERS = ("\ufffd", "\x00")


class SourceDataError(ValueError):
    """Base error for a fail-closed chk1 source-data boundary."""


class InventoryError(SourceDataError):
    """Raised when an inventory entry cannot be bound immutably."""


class StudentPromptSchemaError(SourceDataError):
    """Raised when a legacy student-prompt row no longer matches its schema."""


class CanonicalConflictError(SourceDataError):
    """Raised rather than choosing one of two conflicting canonical records."""


class EvidenceConflictError(SourceDataError):
    """Raised when evidence IDs or factual identities disagree."""


class EvidenceValidationError(SourceDataError):
    """An expected point-in-time exclusion with a machine-readable code."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


class PromptBudgetError(SourceDataError):
    """Raised when final prompt text is invalid or exceeds its hard budget."""


def _canonical_json(value: Any) -> str:
    try:
        return json.dumps(
            value,
            ensure_ascii=False,
            allow_nan=False,
            separators=(",", ":"),
            sort_keys=True,
        )
    except (TypeError, ValueError) as exc:
        raise SourceDataError(f"Value is not canonical finite JSON: {exc}") from exc


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _sha256_text(value: str) -> str:
    return _sha256_bytes(value.encode("utf-8"))


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _require_sha256(value: Any, *, label: str) -> str:
    text = str(value or "").strip().lower()
    if not _SHA256_RE.fullmatch(text):
        raise EvidenceValidationError(
            "invalid_lineage_sha256",
            f"{label} must be a lowercase 64-character SHA-256 digest",
        )
    return text


def _clean_text(value: Any, *, label: str, allow_empty: bool = False) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = text.replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    if not allow_empty and not text:
        raise SourceDataError(f"{label} must be non-empty")
    if any(marker in text for marker in _MOJIBAKE_MARKERS):
        raise SourceDataError(f"{label} contains an encoding anomaly")
    try:
        text.encode("utf-8", errors="strict")
    except UnicodeEncodeError as exc:
        raise SourceDataError(f"{label} is not valid UTF-8 text") from exc
    return text


def _canonical_date(value: Any, *, label: str) -> date:
    text = str(value or "").strip()
    try:
        parsed = date.fromisoformat(text)
    except ValueError as exc:
        raise SourceDataError(f"{label} must be YYYY-MM-DD, got {text!r}") from exc
    if parsed.isoformat() != text:
        raise SourceDataError(f"{label} must be canonical YYYY-MM-DD, got {text!r}")
    return parsed


def _aware_timestamp(value: Any, *, label: str) -> datetime:
    text = str(value or "").strip()
    candidate = text[:-1] + "+00:00" if text.endswith(("Z", "z")) else text
    try:
        parsed = datetime.fromisoformat(candidate)
    except ValueError as exc:
        raise SourceDataError(f"{label} must be an ISO-8601 timestamp") from exc
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise SourceDataError(f"{label} must include a UTC offset")
    return parsed.astimezone(timezone.utc)


def _format_timestamp(value: datetime) -> str:
    return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def classify_inventory_path(path: str | Path) -> str:
    """Classify a legacy path without opening it.

    The important invariant is that any old per-sample target or text capable
    of carrying a target is privileged.  In particular, all legacy llama SFT
    artifacts are explicitly quarantined as required by the chk1 rebuild.
    """

    text = Path(path).as_posix().lower()
    parts = set(Path(text).parts)
    if "processed_llama" in parts and "analysis_sft" in parts:
        return "legacy_privileged"
    if any(token in text for token in ("teacher_prompt", "teacher_response")):
        return "legacy_privileged"
    if "minutes" in text or "merged_response" in text:
        return "legacy_privileged"
    if "student_prompts" in text:
        return "candidate_metadata"
    if "alfred" in text and any(
        token in text for token in ("snapshot", "ledger", "source_evidence", "raw/")
    ):
        return "point_in_time_evidence"
    if "us_data" in parts or "input_sources" in parts:
        return "legacy_unverified"
    return "legacy_audit_only"


def _row_count(path: Path) -> int | None:
    suffix = path.suffix.lower()
    if suffix == ".jsonl":
        with path.open("r", encoding="utf-8") as handle:
            return sum(1 for line in handle if line.strip())
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                count = sum(1 for _ in csv.reader(handle, delimiter=delimiter))
        except (UnicodeDecodeError, csv.Error) as exc:
            raise InventoryError(f"Unable to count rows in {path}: {exc}") from exc
        return max(0, count - 1)
    if suffix == ".json":
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise InventoryError(f"Unable to parse inventory JSON {path}: {exc}") from exc
        return len(payload) if isinstance(payload, list) else 1
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - the project pins openpyxl
            raise InventoryError("openpyxl is required to inventory XLSX rows") from exc
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                return sum(max(0, sheet.max_row - 1) for sheet in workbook.worksheets)
            finally:
                workbook.close()
        except Exception as exc:  # openpyxl exposes several format exceptions
            raise InventoryError(f"Unable to count rows in {path}: {exc}") from exc
    return None


def inventory_file(
    path: str | Path,
    *,
    repo_root: str | Path,
    classification: str | None = None,
) -> dict[str, Any]:
    """Return a content-bound inventory record for one existing regular file."""

    root = Path(repo_root).expanduser().resolve()
    supplied = Path(path).expanduser()
    candidate = supplied if supplied.is_absolute() else root / supplied
    if candidate.is_symlink():
        raise InventoryError(f"Inventory paths may not be symlinks: {candidate}")
    resolved = candidate.resolve()
    try:
        relative = resolved.relative_to(root).as_posix()
    except ValueError as exc:
        raise InventoryError(f"Inventory path escapes repository root: {path}") from exc
    if not resolved.is_file():
        raise InventoryError(f"Inventory path is not a regular file: {resolved}")
    observed_classification = classification or classify_inventory_path(relative)
    if not _CLASSIFICATION_RE.fullmatch(observed_classification):
        raise InventoryError(
            "classification must be a lowercase snake_case identifier, got "
            f"{observed_classification!r}"
        )
    return {
        "path": relative,
        "bytes": resolved.stat().st_size,
        "rows": _row_count(resolved),
        "sha256": _sha256_file(resolved),
        "classification": observed_classification,
    }


def inventory_paths(
    paths: Iterable[str | Path],
    *,
    repo_root: str | Path,
    classifications: Mapping[str, str] | None = None,
) -> list[dict[str, Any]]:
    """Inventory paths deterministically and reject duplicate path aliases."""

    records: dict[str, dict[str, Any]] = {}
    overrides = dict(classifications or {})
    for path in paths:
        provisional = inventory_file(path, repo_root=repo_root)
        relative = provisional["path"]
        classification = overrides.get(relative)
        record = (
            inventory_file(
                path,
                repo_root=repo_root,
                classification=classification,
            )
            if classification is not None
            else provisional
        )
        if relative in records:
            raise InventoryError(f"Duplicate inventory path: {relative}")
        records[relative] = record
    return [records[key] for key in sorted(records)]


def build_file_inventory(
    paths: Iterable[str | Path],
    *,
    repo_root: str | Path,
    classifications: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    """Build a deterministic inventory manifest without modifying any source."""

    files = inventory_paths(
        paths,
        repo_root=repo_root,
        classifications=classifications,
    )
    payload = {
        "schema_version": INVENTORY_SCHEMA_VERSION,
        "file_count": len(files),
        "total_bytes": sum(record["bytes"] for record in files),
        "files": files,
    }
    return {**payload, "payload_sha256": _sha256_text(_canonical_json(payload))}


def _normalise_topic(value: Any) -> str:
    text = _clean_text(value, label="topic")
    return text.replace(" ,", ",")


def split_atomic_topics(topic: Any) -> tuple[str, ...]:
    """Split top-level commas while preserving commas inside parentheses."""

    text = _normalise_topic(topic)
    pieces: list[str] = []
    start = 0
    depth = 0
    for index, character in enumerate(text):
        if character == "(":
            depth += 1
        elif character == ")":
            depth -= 1
            if depth < 0:
                raise StudentPromptSchemaError(f"Unbalanced topic parentheses: {text!r}")
        elif character == "," and depth == 0:
            pieces.append(text[start:index])
            start = index + 1
    if depth != 0:
        raise StudentPromptSchemaError(f"Unbalanced topic parentheses: {text!r}")
    pieces.append(text[start:])
    normalised = tuple(_clean_text(piece, label="atomic_topic") for piece in pieces)
    if len({piece.casefold() for piece in normalised}) != len(normalised):
        raise StudentPromptSchemaError(f"Duplicate atomic topic in {text!r}")
    return normalised


def canonical_key(meeting_date: Any, atomic_topic: Any) -> tuple[str, str]:
    """Return the normalized canonical key used by chk1 and chk2."""

    meeting = _canonical_date(meeting_date, label="meeting_date").isoformat()
    topic = _normalise_topic(atomic_topic)
    if "," in topic:
        raise SourceDataError("canonical atomic_topic must not contain a top-level comma")
    return meeting, topic.casefold()


def stable_sample_id(meeting_date: Any, atomic_topic: Any) -> str:
    meeting, topic_key = canonical_key(meeting_date, atomic_topic)
    digest = _sha256_text(
        _canonical_json({"meeting_date": meeting, "atomic_topic": topic_key})
    )
    return f"chk1-analysis-{meeting}-{digest[:16]}"


def _style_id_from_section(section_name: str) -> str:
    normalized = _clean_text(section_name, label="section_name")
    return f"section-style-{_sha256_text(normalized.casefold())[:16]}"


def _validate_student_prompt_row(row: Mapping[str, Any], *, row_number: int) -> None:
    missing = sorted(STUDENT_PROMPT_REQUIRED_FIELDS - set(row))
    if missing:
        raise StudentPromptSchemaError(
            f"student prompt row {row_number} is missing fields: {missing}"
        )
    sample_id = _clean_text(row.get("sample_id"), label="sample_id")
    split = str(row.get("split") or "").strip()
    if split not in ALLOWED_SPLITS:
        raise StudentPromptSchemaError(f"{sample_id}: invalid split {split!r}")
    _canonical_date(row.get("meeting_date"), label=f"{sample_id}.meeting_date")
    _clean_text(row.get("section_name"), label=f"{sample_id}.section_name")
    split_atomic_topics(row.get("topic"))
    source_row_index = row.get("source_row_index")
    if not isinstance(source_row_index, int) or isinstance(source_row_index, bool):
        raise StudentPromptSchemaError(f"{sample_id}: source_row_index must be an integer")
    prompt = str(row.get("prompt") or "")
    if not prompt:
        raise StudentPromptSchemaError(f"{sample_id}: prompt must be non-empty")
    if str(row.get("prompt_hash") or "").lower() != _sha256_text(prompt):
        raise StudentPromptSchemaError(f"{sample_id}: prompt_hash does not bind prompt")
    if str(row.get("reference_excerpt") or "").strip():
        raise StudentPromptSchemaError(
            f"{sample_id}: student prompt is not reference-free"
        )
    for field in ("quality_flags", "missing_indicators", "source_files"):
        if not isinstance(row.get(field), list):
            raise StudentPromptSchemaError(f"{sample_id}: {field} must be a list")


def build_canonical_samples(
    student_rows: Sequence[Mapping[str, Any]],
    *,
    section_style_by_topic: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    """Project legacy student rows onto safe, unique chk1 sample metadata.

    Exact duplicate metadata are collapsed with their legacy identifiers kept
    only as lineage.  Any disagreement in meeting split or style identity is a
    hard error.  Supplying ``section_style_by_topic`` is the explicit way to
    resolve legacy rows whose section names differ; no row-order tie-break is
    ever used.
    """

    style_map: dict[str, str] | None = None
    if section_style_by_topic is not None:
        style_map = {}
        for raw_topic, raw_style in section_style_by_topic.items():
            topic = _normalise_topic(raw_topic)
            if len(split_atomic_topics(topic)) != 1:
                raise CanonicalConflictError(
                    f"Style map key must be one atomic topic: {raw_topic!r}"
                )
            style = _clean_text(raw_style, label=f"style for {topic}")
            key = topic.casefold()
            if key in style_map and style_map[key] != style:
                raise CanonicalConflictError(f"Conflicting style map entries for {topic}")
            style_map[key] = style

    samples: dict[tuple[str, str], dict[str, Any]] = {}
    meeting_splits: dict[str, str] = {}
    topic_styles: dict[str, str] = {}
    exclusions: list[dict[str, Any]] = []

    for row_number, row in enumerate(student_rows, 1):
        if not isinstance(row, Mapping):
            raise StudentPromptSchemaError(f"student prompt row {row_number} is not an object")
        _validate_student_prompt_row(row, row_number=row_number)
        legacy_sample_id = str(row["sample_id"])
        meeting = str(row["meeting_date"])
        split = str(row["split"])
        prior_split = meeting_splits.setdefault(meeting, split)
        if prior_split != split:
            raise CanonicalConflictError(
                f"Meeting {meeting} appears in both {prior_split!r} and {split!r}"
            )

        section_name = _clean_text(row["section_name"], label="section_name")
        for topic in split_atomic_topics(row["topic"]):
            topic_key = topic.casefold()
            if topic_key in FORBIDDEN_ATOMIC_TOPICS:
                exclusions.append(
                    {
                        "source_sample_id": legacy_sample_id,
                        "meeting_date": meeting,
                        "atomic_topic": topic,
                        "reason": "abnormal_topic",
                    }
                )
                continue
            if style_map is not None:
                style_id = style_map.get(topic_key)
                if style_id is None:
                    raise CanonicalConflictError(
                        f"No explicit section_style_id for atomic topic {topic!r}"
                    )
            else:
                style_id = _style_id_from_section(section_name)

            prior_style = topic_styles.setdefault(topic_key, style_id)
            if prior_style != style_id:
                raise CanonicalConflictError(
                    f"Atomic topic {topic!r} maps to conflicting section styles: "
                    f"{prior_style!r} and {style_id!r}"
                )

            key = canonical_key(meeting, topic)
            candidate = samples.get(key)
            if candidate is None:
                candidate = {
                    "schema_version": CANONICAL_SAMPLE_SCHEMA_VERSION,
                    "sample_id": stable_sample_id(meeting, topic),
                    "canonical_key": {
                        "meeting_date": meeting,
                        "atomic_topic": topic,
                    },
                    "meeting_date": meeting,
                    "atomic_topic": topic,
                    "split": split,
                    "section_style_id": style_id,
                    "legacy_sample_ids": [],
                    "legacy_source_row_indices": [],
                    "legacy_section_names": [],
                }
                samples[key] = candidate
            elif (
                candidate["split"] != split
                or candidate["section_style_id"] != style_id
                or candidate["atomic_topic"] != topic
            ):
                raise CanonicalConflictError(
                    f"Conflicting metadata for canonical key {meeting!r} + {topic!r}"
                )
            candidate["legacy_sample_ids"].append(legacy_sample_id)
            candidate["legacy_source_row_indices"].append(int(row["source_row_index"]))
            candidate["legacy_section_names"].append(section_name)

    ordered_samples = []
    for key in sorted(samples):
        sample = samples[key]
        for field in (
            "legacy_sample_ids",
            "legacy_source_row_indices",
            "legacy_section_names",
        ):
            sample[field] = sorted(set(sample[field]))
        ordered_samples.append(sample)

    payload = {
        "schema_version": CANONICAL_BATCH_SCHEMA_VERSION,
        "canonical_key_fields": ["meeting_date", "atomic_topic"],
        "samples": ordered_samples,
        "exclusions": sorted(
            exclusions,
            key=lambda item: (
                item["meeting_date"],
                item["source_sample_id"],
                item["atomic_topic"],
            ),
        ),
    }
    return {**payload, "payload_sha256": _sha256_text(_canonical_json(payload))}


def _evidence_field(row: Mapping[str, Any], name: str) -> Any:
    value = row.get(name)
    if value not in (None, ""):
        return value
    lineage = row.get("lineage")
    if isinstance(lineage, Mapping):
        return lineage.get(name)
    return value


def _normalise_numeric_value(value: Any, *, evidence_id: str) -> str:
    if isinstance(value, bool) or value is None:
        raise EvidenceValidationError(
            "invalid_value", f"{evidence_id}: value must be numeric or a factual string"
        )
    if isinstance(value, float) and not math.isfinite(value):
        raise EvidenceValidationError("invalid_value", f"{evidence_id}: value is not finite")
    text = str(value).strip()
    if not text or text.lower() in {"nan", "inf", "+inf", "-inf", "null", "none"}:
        raise EvidenceValidationError("invalid_value", f"{evidence_id}: value is empty/non-finite")
    # Preserve exact decimal strings when possible.  Descriptive trend values
    # remain strings and must still be linked to their evidence/formula.
    try:
        decimal = Decimal(text)
    except InvalidOperation:
        return _clean_text(text, label=f"{evidence_id}.value")
    if not decimal.is_finite():
        raise EvidenceValidationError("invalid_value", f"{evidence_id}: value is not finite")
    return text


def _evidence_error(code: str, message: str) -> EvidenceValidationError:
    return EvidenceValidationError(code, message)


def validate_point_in_time_evidence(
    row: Mapping[str, Any],
    *,
    meeting_date: str,
    atomic_topic: str,
    cutoff_ts: str,
) -> dict[str, Any]:
    """Validate and normalize one fact candidate against its information cutoff.

    Macro observations require the repository's canonical ALFRED D-1 vintage
    attestation plus raw/request/manifest hashes.  A boolean such as
    ``vintage_verified=true`` is deliberately not accepted as provenance.
    """

    if not isinstance(row, Mapping):
        raise _evidence_error("invalid_evidence", "Evidence row must be an object")
    meeting = _canonical_date(meeting_date, label="meeting_date")
    topic = _normalise_topic(atomic_topic)
    cutoff = _aware_timestamp(cutoff_ts, label="cutoff_ts")
    evidence_id = _clean_text(row.get("evidence_id"), label="evidence_id")

    row_meeting = row.get("meeting_date")
    if row_meeting not in (None, "") and _canonical_date(
        row_meeting, label=f"{evidence_id}.meeting_date"
    ) != meeting:
        raise _evidence_error("meeting_mismatch", f"{evidence_id}: meeting_date mismatch")
    row_topic = row.get("atomic_topic", row.get("indicator"))
    if row_topic not in (None, "") and _normalise_topic(row_topic).casefold() != topic.casefold():
        raise _evidence_error("topic_mismatch", f"{evidence_id}: atomic_topic mismatch")
    bound_cutoff = row.get("cutoff_ts")
    if bound_cutoff not in (None, "") and _aware_timestamp(
        bound_cutoff, label=f"{evidence_id}.cutoff_ts"
    ) != cutoff:
        raise _evidence_error("cutoff_mismatch", f"{evidence_id}: cutoff_ts mismatch")

    source_id = _clean_text(_evidence_field(row, "source_id"), label="source_id")
    source_sha256 = _require_sha256(
        _evidence_field(row, "source_sha256"), label=f"{evidence_id}.source_sha256"
    )
    source_kind = str(row.get("source_kind") or "macro").strip().lower()
    series_id = _clean_text(
        row.get("series_id", source_id), label=f"{evidence_id}.series_id"
    )
    metric = _clean_text(
        row.get("metric", row.get("title", series_id)), label=f"{evidence_id}.metric"
    )
    units = _clean_text(row.get("units"), label=f"{evidence_id}.units")
    fact_kind = str(row.get("fact_kind") or "recent_observation").strip().lower()
    if fact_kind not in ALLOWED_FACT_KINDS:
        raise _evidence_error("invalid_fact_kind", f"{evidence_id}: unsupported fact_kind")
    observation = _canonical_date(
        row.get("observation_date"), label=f"{evidence_id}.observation_date"
    )
    if observation > cutoff.date():
        raise _evidence_error(
            "observation_after_cutoff", f"{evidence_id}: observation is after cutoff"
        )
    release: datetime | None = None
    release_raw = row.get("release_ts")
    if release_raw not in (None, ""):
        try:
            release = _aware_timestamp(
                release_raw, label=f"{evidence_id}.release_ts"
            )
        except SourceDataError as exc:
            raise _evidence_error("invalid_release_ts", str(exc)) from exc
        if release > cutoff:
            raise _evidence_error(
                "release_after_cutoff", f"{evidence_id}: release is after cutoff"
            )

    raw_sha256 = _require_sha256(
        _evidence_field(row, "raw_sha256"), label=f"{evidence_id}.raw_sha256"
    )
    source_interface = _clean_text(
        row.get("source_interface"), label=f"{evidence_id}.source_interface"
    )
    evidence_type = _clean_text(
        row.get("availability_evidence_type"),
        label=f"{evidence_id}.availability_evidence_type",
    )

    requested_vintage_date: str | None = None
    information_as_of_date: str | None = None
    availability_as_of_date: str | None = None
    request_id: str | None = None
    snapshot_manifest_payload_sha256: str | None = None
    availability_upper_bound: datetime | None = None
    availability_basis: str | None = None

    if source_kind == "macro":
        if (
            source_interface != ALFRED_SOURCE_INTERFACE
            or evidence_type != ALFRED_AVAILABILITY_EVIDENCE
        ):
            raise _evidence_error(
                "unverified_vintage",
                f"{evidence_id}: macro evidence lacks canonical ALFRED vintage provenance",
            )
        try:
            vintage = _canonical_date(
                row.get("requested_vintage_date", row.get("vintage_date")),
                label=f"{evidence_id}.requested_vintage_date",
            )
            information_as_of = _canonical_date(
                row.get("information_as_of_date"),
                label=f"{evidence_id}.information_as_of_date",
            )
            availability_as_of = _canonical_date(
                row.get("availability_as_of_date"),
                label=f"{evidence_id}.availability_as_of_date",
            )
        except SourceDataError as exc:
            raise _evidence_error("missing_vintage_provenance", str(exc)) from exc
        expected_vintage = meeting - timedelta(days=1)
        if not (vintage == information_as_of == availability_as_of == expected_vintage):
            raise _evidence_error(
                "vintage_cutoff_mismatch",
                f"{evidence_id}: ALFRED vintage must equal meeting_date - 1 day",
            )
        if vintage > cutoff.date():
            raise _evidence_error("vintage_after_cutoff", f"{evidence_id}: vintage is after cutoff")
        request_id = _require_sha256(
            _evidence_field(row, "request_id"), label=f"{evidence_id}.request_id"
        )
        snapshot_manifest_payload_sha256 = _require_sha256(
            _evidence_field(row, "snapshot_manifest_payload_sha256"),
            label=f"{evidence_id}.snapshot_manifest_payload_sha256",
        )
        requested_vintage_date = vintage.isoformat()
        information_as_of_date = information_as_of.isoformat()
        availability_as_of_date = availability_as_of.isoformat()
        upper_bound_raw = row.get("availability_upper_bound_ts")
        if upper_bound_raw not in (None, ""):
            try:
                availability_upper_bound = _aware_timestamp(
                    upper_bound_raw,
                    label=f"{evidence_id}.availability_upper_bound_ts",
                )
            except SourceDataError as exc:
                raise _evidence_error(
                    "invalid_availability_upper_bound", str(exc)
                ) from exc
            expected_upper_bound = datetime.combine(
                expected_vintage,
                datetime.max.time().replace(microsecond=0),
                tzinfo=timezone.utc,
            )
            if availability_upper_bound != expected_upper_bound:
                raise _evidence_error(
                    "availability_upper_bound_mismatch",
                    f"{evidence_id}: sealed ALFRED upper bound must be D-1 23:59:59Z",
                )
            if availability_upper_bound > cutoff:
                raise _evidence_error(
                    "availability_after_cutoff",
                    f"{evidence_id}: availability upper bound is after cutoff",
                )
        if release is not None:
            availability_basis = "actual_release_ts"
        elif availability_upper_bound is not None:
            # This is a conservative upper bound proven by the sealed D-1
            # snapshot.  It is not represented as the observation's actual
            # release timestamp.
            availability_basis = "sealed_alfred_d1_upper_bound"
        else:
            raise _evidence_error(
                "missing_availability_proof",
                f"{evidence_id}: provide actual release_ts or the sealed D-1 "
                "availability_upper_bound_ts",
            )
    elif source_kind == "market":
        if release is None:
            raise _evidence_error(
                "missing_release_ts",
                f"{evidence_id}: market evidence requires an actual release_ts",
            )
        if evidence_type != "market_close_snapshot":
            raise _evidence_error(
                "unverified_market_close", f"{evidence_id}: missing market-close provenance"
            )
        market_close_raw = row.get("market_close_ts", row.get("observation_ts"))
        try:
            market_close = _aware_timestamp(
                market_close_raw, label=f"{evidence_id}.market_close_ts"
            )
        except SourceDataError as exc:
            raise _evidence_error("missing_market_close_ts", str(exc)) from exc
        if market_close > cutoff:
            raise _evidence_error(
                "market_close_after_cutoff", f"{evidence_id}: market close is after cutoff"
            )
        availability_basis = "actual_release_ts"
    elif source_kind == "policy":
        if release is None:
            raise _evidence_error(
                "missing_release_ts",
                f"{evidence_id}: policy evidence requires an actual release_ts",
            )
        if evidence_type != "official_policy_release":
            raise _evidence_error(
                "unverified_policy_release", f"{evidence_id}: missing official policy release"
            )
        try:
            effective = _aware_timestamp(
                row.get("effective_ts"), label=f"{evidence_id}.effective_ts"
            )
        except SourceDataError as exc:
            raise _evidence_error("missing_policy_effective_ts", str(exc)) from exc
        if effective >= cutoff:
            raise _evidence_error(
                "policy_not_prior_to_cutoff",
                f"{evidence_id}: policy value is not effective before cutoff",
            )
        if fact_kind != "prior_target_range":
            raise _evidence_error(
                "invalid_policy_fact", f"{evidence_id}: policy evidence must be prior_target_range"
            )
        availability_basis = "actual_release_ts"
    else:
        raise _evidence_error(
            "unsupported_source_kind", f"{evidence_id}: unsupported source_kind {source_kind!r}"
        )

    formula = str(row.get("formula") or "").strip()
    operands_raw = row.get("operand_evidence_ids", [])
    if not isinstance(operands_raw, list) or any(
        not isinstance(value, str) or not value.strip() for value in operands_raw
    ):
        raise _evidence_error(
            "invalid_formula_lineage",
            f"{evidence_id}: operand_evidence_ids must be a list of non-empty strings",
        )
    operands = list(dict.fromkeys(value.strip() for value in operands_raw))
    if evidence_id in operands:
        raise _evidence_error(
            "invalid_formula_lineage", f"{evidence_id}: formula may not depend on itself"
        )
    if fact_kind in DERIVED_FACT_KINDS and (not formula or not operands):
        raise _evidence_error(
            "missing_formula_lineage",
            f"{evidence_id}: derived fact requires formula and operand evidence IDs",
        )

    registry_sha = _evidence_field(row, "registry_sha256")
    registry_sha256 = (
        _require_sha256(registry_sha, label=f"{evidence_id}.registry_sha256")
        if registry_sha not in (None, "")
        else None
    )
    normalized = {
        "evidence_id": evidence_id,
        "meeting_date": meeting.isoformat(),
        "atomic_topic": topic,
        "source_kind": source_kind,
        "source_id": source_id,
        "series_id": series_id,
        "metric": metric,
        "fact_kind": fact_kind,
        "value": _normalise_numeric_value(row.get("value"), evidence_id=evidence_id),
        "units": units,
        "observation_date": observation.isoformat(),
        "release_ts": _format_timestamp(release) if release is not None else None,
        "availability_upper_bound_ts": (
            _format_timestamp(availability_upper_bound)
            if availability_upper_bound is not None
            else None
        ),
        "availability_basis": availability_basis,
        "cutoff_ts": _format_timestamp(cutoff),
        "requested_vintage_date": requested_vintage_date,
        "information_as_of_date": information_as_of_date,
        "availability_as_of_date": availability_as_of_date,
        "availability_evidence_type": evidence_type,
        "source_interface": source_interface,
        "formula": formula,
        "operand_evidence_ids": operands,
        "lineage": {
            "source_sha256": source_sha256,
            "raw_sha256": raw_sha256,
            "request_id": request_id,
            "snapshot_manifest_payload_sha256": snapshot_manifest_payload_sha256,
            "registry_sha256": registry_sha256,
            "availability_basis": availability_basis,
        },
    }
    normalized["evidence_sha256"] = _sha256_text(_canonical_json(normalized))
    return normalized


def _exclusion(row: Mapping[str, Any], *, reason: str, detail: str) -> dict[str, Any]:
    try:
        candidate_sha = _sha256_text(_canonical_json(dict(row)))
    except SourceDataError:
        candidate_sha = None
    return {
        "evidence_id": str(row.get("evidence_id") or ""),
        "source_id": str(_evidence_field(row, "source_id") or ""),
        "reason": reason,
        "detail": detail,
        "candidate_sha256": candidate_sha,
    }


def build_fact_card(
    *,
    meeting_date: str,
    atomic_topic: str,
    cutoff_ts: str,
    evidence_rows: Sequence[Mapping[str, Any]],
    max_recent_observations: int = 6,
) -> dict[str, Any]:
    """Build a compact fact card, excluding any unprovable evidence.

    Invalid individual evidence becomes an auditable exclusion.  Conflicting
    evidence is different: the function raises and emits no fact card, so a
    caller cannot accidentally select a preferred value by input order.
    """

    meeting = _canonical_date(meeting_date, label="meeting_date").isoformat()
    topic = _normalise_topic(atomic_topic)
    cutoff = _format_timestamp(_aware_timestamp(cutoff_ts, label="cutoff_ts"))
    if topic.casefold() in FORBIDDEN_ATOMIC_TOPICS:
        return {
            "schema_version": FACT_CARD_BUILD_SCHEMA_VERSION,
            "status": "excluded",
            "fact_card": None,
            "exclusions": [
                {
                    "evidence_id": "",
                    "source_id": "",
                    "reason": "abnormal_topic",
                    "detail": f"atomic topic {topic!r} is not eligible",
                    "candidate_sha256": None,
                }
            ],
        }
    if (
        not isinstance(max_recent_observations, int)
        or isinstance(max_recent_observations, bool)
        or max_recent_observations <= 0
    ):
        raise SourceDataError("max_recent_observations must be a positive integer")

    accepted: dict[str, dict[str, Any]] = {}
    identities: dict[tuple[str, str, str], dict[str, Any]] = {}
    exclusions: list[dict[str, Any]] = []
    for raw in evidence_rows:
        if not isinstance(raw, Mapping):
            exclusions.append(
                {
                    "evidence_id": "",
                    "source_id": "",
                    "reason": "invalid_evidence",
                    "detail": "Evidence row must be an object",
                    "candidate_sha256": None,
                }
            )
            continue
        try:
            normalized = validate_point_in_time_evidence(
                raw,
                meeting_date=meeting,
                atomic_topic=topic,
                cutoff_ts=cutoff,
            )
        except EvidenceValidationError as exc:
            exclusions.append(_exclusion(raw, reason=exc.code, detail=str(exc)))
            continue
        evidence_id = normalized["evidence_id"]
        prior = accepted.get(evidence_id)
        if prior is not None:
            if prior != normalized:
                raise EvidenceConflictError(
                    f"Evidence ID {evidence_id!r} binds different payloads"
                )
            continue
        identity = (
            normalized["series_id"].casefold(),
            normalized["observation_date"],
            normalized["fact_kind"],
        )
        prior_identity = identities.get(identity)
        if prior_identity is not None and (
            prior_identity["value"] != normalized["value"]
            or prior_identity["units"] != normalized["units"]
        ):
            raise EvidenceConflictError(
                "Conflicting values for factual identity "
                f"{identity}: {prior_identity['evidence_id']!r} vs {evidence_id!r}"
            )
        identities[identity] = normalized
        accepted[evidence_id] = normalized

    # Derived facts may only depend on evidence that survived every point-in-
    # time check.  Iterate because excluding one derived fact can invalidate a
    # second-order derivation.
    changed = True
    while changed:
        changed = False
        for evidence_id, normalized in list(accepted.items()):
            missing_operands = sorted(
                set(normalized["operand_evidence_ids"]) - set(accepted)
            )
            if missing_operands:
                exclusions.append(
                    {
                        "evidence_id": evidence_id,
                        "source_id": normalized["source_id"],
                        "reason": "missing_formula_operand",
                        "detail": f"Missing accepted operands: {missing_operands}",
                        "candidate_sha256": normalized["evidence_sha256"],
                    }
                )
                del accepted[evidence_id]
                changed = True

    recent_by_series: dict[str, list[dict[str, Any]]] = defaultdict(list)
    derived: list[dict[str, Any]] = []
    for normalized in accepted.values():
        if normalized["fact_kind"] in {"recent_observation", "latest"}:
            recent_by_series[normalized["series_id"]].append(normalized)
        else:
            derived.append(normalized)

    selected: list[dict[str, Any]] = list(derived)
    for series_id in sorted(recent_by_series):
        observations = sorted(
            recent_by_series[series_id],
            key=lambda item: (item["observation_date"], item["evidence_id"]),
        )
        dropped = observations[:-max_recent_observations]
        observations = observations[-max_recent_observations:]
        for normalized in dropped:
            exclusions.append(
                {
                    "evidence_id": normalized["evidence_id"],
                    "source_id": normalized["source_id"],
                    "reason": "outside_fact_card_recent_limit",
                    "detail": (
                        f"Only the latest {max_recent_observations} observations "
                        f"for {series_id} are retained"
                    ),
                    "candidate_sha256": normalized["evidence_sha256"],
                }
            )
        if observations:
            for normalized in observations[:-1]:
                normalized = dict(normalized)
                normalized["fact_kind"] = "recent_observation"
                selected.append(normalized)
            latest = dict(observations[-1])
            latest["fact_kind"] = "latest"
            selected.append(latest)

    selected.sort(
        key=lambda item: (
            item["series_id"].casefold(),
            item["observation_date"],
            item["fact_kind"],
            item["evidence_id"],
        )
    )
    exclusions.sort(
        key=lambda item: (
            item.get("evidence_id", ""),
            item.get("reason", ""),
            item.get("candidate_sha256") or "",
        )
    )
    if not selected:
        return {
            "schema_version": FACT_CARD_BUILD_SCHEMA_VERSION,
            "status": "excluded",
            "fact_card": None,
            "exclusions": exclusions
            or [
                {
                    "evidence_id": "",
                    "source_id": "",
                    "reason": "empty_fact_card",
                    "detail": "No point-in-time evidence survived",
                    "candidate_sha256": None,
                }
            ],
        }

    evidence_items: list[dict[str, Any]] = []
    lineage: list[dict[str, Any]] = []
    for item in selected:
        evidence_items.append(
            {
                "evidence_id": item["evidence_id"],
                "series_id": item["series_id"],
                "metric": item["metric"],
                "fact_kind": item["fact_kind"],
                "value": item["value"],
                "units": item["units"],
                "observation_date": item["observation_date"],
                "formula": item["formula"],
                "operand_evidence_ids": item["operand_evidence_ids"],
                "release_ts": item["release_ts"],
                "availability_upper_bound_ts": item[
                    "availability_upper_bound_ts"
                ],
                "availability_basis": item["availability_basis"],
                "requested_vintage_date": item["requested_vintage_date"],
                "cutoff_ts": item["cutoff_ts"],
                "source_sha256": item["lineage"]["source_sha256"],
            }
        )
        lineage.append(
            {
                "evidence_id": item["evidence_id"],
                "evidence_sha256": item["evidence_sha256"],
                "source_id": item["source_id"],
                "source_kind": item["source_kind"],
                "source_interface": item["source_interface"],
                "availability_evidence_type": item[
                    "availability_evidence_type"
                ],
                "source_sha256": item["lineage"]["source_sha256"],
                "raw_sha256": item["lineage"]["raw_sha256"],
                "request_id": item["lineage"]["request_id"],
                "snapshot_manifest_payload_sha256": item["lineage"][
                    "snapshot_manifest_payload_sha256"
                ],
                "registry_sha256": item["lineage"]["registry_sha256"],
                "release_ts": item["release_ts"],
                "availability_upper_bound_ts": item[
                    "availability_upper_bound_ts"
                ],
                "availability_basis": item["availability_basis"],
                "requested_vintage_date": item["requested_vintage_date"],
                "information_as_of_date": item["information_as_of_date"],
                "availability_as_of_date": item["availability_as_of_date"],
                "cutoff_ts": item["cutoff_ts"],
            }
        )

    fact_card = {
        "schema_version": FACT_CARD_SCHEMA_VERSION,
        "sample_id": stable_sample_id(meeting, topic),
        "canonical_key": {"meeting_date": meeting, "atomic_topic": topic},
        "meeting_date": meeting,
        "atomic_topic": topic,
        "cutoff_ts": cutoff,
        # ``evidence`` is the canonical verifier-facing fact-card payload.
        # Keeping one key avoids duplicating token-bearing content in prompts.
        "evidence": evidence_items,
        "evidence_lineage": lineage,
    }
    fact_card["fact_card_sha256"] = _sha256_text(_canonical_json(fact_card))
    return {
        "schema_version": FACT_CARD_BUILD_SCHEMA_VERSION,
        "status": "ready",
        "fact_card": fact_card,
        "exclusions": exclusions,
    }


def evidence_from_loo_ledger(
    ledger_row: Mapping[str, Any],
    source_evidence_row: Mapping[str, Any],
    *,
    atomic_topic: str | None = None,
    release_ts_by_observation: Mapping[tuple[str, str], str] | None = None,
    cutoff_ts: str | None = None,
) -> list[dict[str, Any]]:
    """Project a canonical ALFRED/LOO row into chk1 evidence candidates.

    The adapter replays the existing source-payload and selected-observation
    hashes.  If an exact release timestamp is unavailable, it records D-1
    23:59:59Z as a conservative ``availability_upper_bound_ts`` proven by the
    sealed ALFRED vintage.  It never labels that bound as the actual release.
    """

    if not isinstance(ledger_row, Mapping) or not isinstance(source_evidence_row, Mapping):
        raise EvidenceConflictError("LOO ledger and source evidence must be objects")
    if ledger_row.get("schema_version") != "canonical-loo-indicator-input-v2":
        raise EvidenceConflictError("Unsupported LOO ledger row schema")
    if source_evidence_row.get("schema_version") != "loo-indicator-source-evidence-v1":
        raise EvidenceConflictError("Unsupported LOO source-evidence schema")
    for field in ("sample_id", "source_id", "source_sha256", "information_as_of_date"):
        if ledger_row.get(field) != source_evidence_row.get(field):
            raise EvidenceConflictError(f"LOO ledger/evidence binding mismatch: {field}")
    payload = ledger_row.get("source_payload")
    if not isinstance(payload, Mapping):
        raise EvidenceConflictError("LOO ledger source_payload must be an object")
    payload_sha = _sha256_text(_canonical_json(payload))
    if payload_sha != source_evidence_row.get("source_payload_sha256"):
        raise EvidenceConflictError("LOO source_payload_sha256 mismatch")
    evidence_binding = {
        key: value
        for key, value in source_evidence_row.items()
        if key != "source_sha256"
    }
    if _sha256_text(_canonical_json(evidence_binding)) != source_evidence_row.get(
        "source_sha256"
    ):
        raise EvidenceConflictError("LOO source evidence digest mismatch")

    requests_raw = source_evidence_row.get("requests")
    if not isinstance(requests_raw, list):
        raise EvidenceConflictError("LOO source evidence requests must be a list")
    requests: dict[str, Mapping[str, Any]] = {}
    for request in requests_raw:
        if not isinstance(request, Mapping):
            raise EvidenceConflictError("LOO request evidence must be an object")
        source_key = str(request.get("source_key") or "")
        if not source_key or source_key in requests:
            raise EvidenceConflictError(f"Duplicate/empty LOO source_key {source_key!r}")
        requests[source_key] = request

    meeting = _canonical_date(ledger_row.get("meeting_date"), label="LOO meeting_date")
    ledger_topic = _normalise_topic(ledger_row.get("indicator"))
    topic = _normalise_topic(atomic_topic or ledger_topic)
    # The canonical LOO ledger uses registry slugs such as ``Bank-Capital``
    # while chk1 freezes human-readable atomic topics such as ``Bank Capital``.
    # Permit only that punctuation-only projection; an actual topic change is
    # a provenance mismatch and must fail closed.
    def topic_identity(value: str) -> str:
        return re.sub(r"[^a-z0-9]+", "", value.casefold())
    if topic_identity(topic) != topic_identity(ledger_topic):
        raise EvidenceConflictError(
            "LOO ledger indicator does not match the requested atomic topic"
        )
    effective_cutoff = cutoff_ts or str(ledger_row.get("meeting_timestamp") or "")
    _aware_timestamp(effective_cutoff, label="LOO cutoff_ts")
    release_lookup = dict(release_ts_by_observation or {})
    information_as_of = _canonical_date(
        ledger_row.get("information_as_of_date"),
        label="LOO information_as_of_date",
    )
    availability_upper_bound_ts = datetime.combine(
        information_as_of,
        datetime.max.time().replace(microsecond=0),
        tzinfo=timezone.utc,
    ).isoformat().replace("+00:00", "Z")
    series_rows = payload.get("series")
    if not isinstance(series_rows, list) or not series_rows:
        raise EvidenceConflictError("LOO source_payload.series must be non-empty")

    candidates: list[dict[str, Any]] = []
    for series in series_rows:
        if not isinstance(series, Mapping):
            raise EvidenceConflictError("LOO series payload must be an object")
        source_key = str(series.get("source_key") or "")
        request = requests.get(source_key)
        if request is None:
            raise EvidenceConflictError(f"Missing LOO request evidence for {source_key!r}")
        observations = series.get("observations")
        if not isinstance(observations, list) or not observations:
            raise EvidenceConflictError(f"LOO series {source_key!r} has no observations")
        if _sha256_text(_canonical_json(observations)) != request.get(
            "selected_observation_sha256"
        ):
            raise EvidenceConflictError(
                f"LOO selected observation digest mismatch for {source_key!r}"
            )
        series_id = _clean_text(series.get("series_id"), label="LOO series_id")
        for observation in observations:
            if not isinstance(observation, Mapping):
                raise EvidenceConflictError("LOO observation must be an object")
            observation_date = _canonical_date(
                observation.get("date"), label="LOO observation date"
            ).isoformat()
            descriptor = {
                "meeting_date": meeting.isoformat(),
                "atomic_topic": topic.casefold(),
                "series_id": series_id,
                "observation_date": observation_date,
                "source_sha256": ledger_row["source_sha256"],
            }
            evidence_id = f"ev-{_sha256_text(_canonical_json(descriptor))[:24]}"
            candidates.append(
                {
                    "evidence_id": evidence_id,
                    "meeting_date": meeting.isoformat(),
                    "atomic_topic": topic,
                    "source_kind": "macro",
                    "source_id": ledger_row["source_id"],
                    "source_sha256": ledger_row["source_sha256"],
                    "series_id": series_id,
                    "metric": series.get("title", series_id),
                    "fact_kind": "recent_observation",
                    "value": observation.get("value"),
                    "units": series.get("units"),
                    "observation_date": observation_date,
                    "release_ts": release_lookup.get(
                        (series_id, observation_date)
                    ),
                    "availability_upper_bound_ts": availability_upper_bound_ts,
                    "cutoff_ts": effective_cutoff,
                    "requested_vintage_date": series.get(
                        "requested_vintage_date",
                        ledger_row.get("requested_vintage_date"),
                    ),
                    "information_as_of_date": ledger_row.get(
                        "information_as_of_date"
                    ),
                    "availability_as_of_date": series.get(
                        "availability_as_of_date",
                        ledger_row.get("availability_as_of_date"),
                    ),
                    "availability_evidence_type": ledger_row.get(
                        "availability_evidence_type"
                    ),
                    "source_interface": ledger_row.get("source_interface"),
                    "lineage": {
                        "raw_sha256": request.get("raw_sha256"),
                        "request_id": request.get("request_id"),
                        "snapshot_manifest_payload_sha256": source_evidence_row.get(
                            "snapshot_manifest_payload_sha256"
                        ),
                        "registry_sha256": source_evidence_row.get(
                            "registry_sha256"
                        ),
                    },
                }
            )
    return sorted(candidates, key=lambda item: item["evidence_id"])


def validate_prompt_token_budget(
    prompt: str,
    *,
    token_counter: Callable[[str], int],
    max_tokens: int = PROMPT_TOKEN_LIMIT,
) -> dict[str, Any]:
    """Validate a final prompt with an injected tokenizer/counting boundary.

    ``token_counter`` must count the exact text that will be supplied to the
    model (including any already-rendered system/chat template).  The function
    has no truncation branch: an overflow is always an error.
    """

    if not isinstance(prompt, str) or not prompt.strip():
        raise PromptBudgetError("Prompt must be non-empty text")
    if any(marker in prompt for marker in _MOJIBAKE_MARKERS):
        raise PromptBudgetError("Prompt contains an encoding anomaly")
    try:
        prompt.encode("utf-8", errors="strict")
    except UnicodeEncodeError as exc:
        raise PromptBudgetError("Prompt is not valid UTF-8 text") from exc
    if (
        not isinstance(max_tokens, int)
        or isinstance(max_tokens, bool)
        or max_tokens <= 0
    ):
        raise PromptBudgetError("max_tokens must be a positive integer")
    if not callable(token_counter):
        raise PromptBudgetError("token_counter must be callable")
    try:
        token_count = token_counter(prompt)
    except Exception as exc:
        raise PromptBudgetError(f"Token counter failed: {exc}") from exc
    if (
        not isinstance(token_count, int)
        or isinstance(token_count, bool)
        or token_count < 0
    ):
        raise PromptBudgetError("Token counter must return a non-negative integer")
    if token_count > max_tokens:
        raise PromptBudgetError(
            f"Prompt has {token_count} tokens, exceeding hard limit {max_tokens}; "
            "truncation is forbidden"
        )
    return {
        "schema_version": PROMPT_BUDGET_SCHEMA_VERSION,
        "prompt_sha256": _sha256_text(prompt),
        "token_count": token_count,
        "max_tokens": max_tokens,
        "overflow_policy": "error",
        "truncated": False,
    }


def make_huggingface_token_counter(
    tokenizer: Any,
    *,
    add_special_tokens: bool = True,
) -> Callable[[str], int]:
    """Adapt a Hugging Face-compatible tokenizer to the external token gate."""

    if not hasattr(tokenizer, "encode") or not callable(tokenizer.encode):
        raise PromptBudgetError("tokenizer must expose a callable encode method")

    def count(text: str) -> int:
        token_ids = tokenizer.encode(text, add_special_tokens=add_special_tokens)
        if not isinstance(token_ids, Sequence):
            raise TypeError("tokenizer.encode must return a token sequence")
        return len(token_ids)

    return count


__all__ = [
    "ALFRED_AVAILABILITY_EVIDENCE",
    "ALFRED_SOURCE_INTERFACE",
    "CANONICAL_SAMPLE_SCHEMA_VERSION",
    "CanonicalConflictError",
    "EvidenceConflictError",
    "EvidenceValidationError",
    "InventoryError",
    "PROMPT_TOKEN_LIMIT",
    "PromptBudgetError",
    "STUDENT_PROMPT_REQUIRED_FIELDS",
    "SourceDataError",
    "StudentPromptSchemaError",
    "build_canonical_samples",
    "build_fact_card",
    "build_file_inventory",
    "canonical_key",
    "classify_inventory_path",
    "evidence_from_loo_ledger",
    "inventory_file",
    "inventory_paths",
    "make_huggingface_token_counter",
    "split_atomic_topics",
    "stable_sample_id",
    "validate_point_in_time_evidence",
    "validate_prompt_token_budget",
]
